import os
import openpyxl
from datetime import datetime
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib import messages
from django.conf import settings
from django.urls import reverse
from django.http import HttpResponse, FileResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.decorators.csrf import csrf_exempt
from django.core.files.storage import default_storage
from django.core.files.base import ContentFile
from PIL import Image
import qrcode
import requests
import logging

from .models import Device, QRCode, Computer, Printer, Employee, Department, DeviceType

logger = logging.getLogger(__name__)


@staff_member_required
def generate_qr_with_path(request):
    """
    Представление для генерации QR-кодов с выбором пути (для админки).
    """
    if request.method == 'POST':
        path = request.POST.get('path', '').strip()
        ids = request.POST.get('ids', '')
        device_ids = [int(id) for id in ids.split(',') if id.isdigit()]
        devices = Device.objects.filter(id__in=device_ids)

        if path:
            target_dir = os.path.join(settings.MEDIA_ROOT, path)
        else:
            target_dir = os.path.join(settings.MEDIA_ROOT, 'qrcodes')

        os.makedirs(target_dir, exist_ok=True)

        generated = 0
        for device in devices:
            qr, created = QRCode.objects.get_or_create(device=device)
            try:
                from ..qr_generator.utils import generate_qr_image_for_device
                generate_qr_image_for_device(device, target_dir=target_dir)
                generated += 1
            except Exception as e:
                messages.error(request, f"Ошибка для устройства {device.inventory_number}: {e}")

        messages.success(request, f"Сгенерировано {generated} QR-кодов в папку {target_dir}")
        return redirect('admin:core_device_changelist')

    ids = request.GET.get('ids', '')
    return render(request, 'admin/qr_path_form.html', {'ids': ids})


@staff_member_required
def assign_qr_to_device(request):
    """
    Представление для привязки выбранных QR-кодов к устройству.
    """
    if request.method == 'POST':
        qr_ids = request.POST.get('qr_ids', '')
        device_id = request.POST.get('device_id')
        qr_list = QRCode.objects.filter(id__in=[int(x) for x in qr_ids.split(',') if x.isdigit()])
        device = get_object_or_404(Device, id=device_id)

        for qr in qr_list:
            try:
                existing_qr = device.qr_code
            except QRCode.DoesNotExist:
                existing_qr = None

            # Привязываем только к свободной технике
            if existing_qr and existing_qr != qr:
                messages.error(
                    request,
                    f"❌ Устройство {device.inventory_number} уже имеет QR-код (ID {existing_qr.id})."
                )
                return redirect('admin:core_qrcode_changelist')

            qr.device = device
            qr.save()
            qr.generate_image()

        messages.success(request, f"✅ Привязано {len(qr_list)} QR-кодов к устройству {device}")
        return redirect('admin:core_qrcode_changelist')

    qr_ids = request.GET.get('qr_ids', '')
    devices = Device.objects.filter(qr_code__isnull=True).order_by('inventory_number')
    return render(request, 'admin/assign_qr_form.html', {
        'qr_ids': qr_ids,
        'devices': devices
    })


@staff_member_required
def computer_list_partial(request):
    """
    Возвращает отрендеренный HTML-фрагмент со списком компьютеров.
    Используется для динамического обновления страницы через AJAX/HTMX.
    """
    computers = Computer.objects.all().order_by('inventory_number')
    return render(request, 'admin/computer_list_partial.html', {'computers': computers})


@require_GET
def generate_qr_api(request):
    """
    API для генерации QR-кода из произвольных данных.
    Параметры GET:
        data - текст для кодирования (обязательно)
        size - размер стороны в пикселях (по умолчанию 200)
        fill_color - цвет QR-кода (по умолчанию black)
        back_color - цвет фона (по умолчанию white)
    Возвращает PNG изображение.
    """
    data = request.GET.get('data')
    if not data:
        return HttpResponse('Missing data parameter', status=400)

    try:
        size = int(request.GET.get('size', 200))
    except ValueError:
        size = 200

    fill_color = request.GET.get('fill_color', 'black')
    back_color = request.GET.get('back_color', 'white')

    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_L,
        box_size=10,
        border=4,
    )
    qr.add_data(data)
    qr.make(fit=True)
    img = qr.make_image(fill_color=fill_color, back_color=back_color).convert('RGB')

    if size and size != img.size[0]:
        img = img.resize((size, size), Image.Resampling.LANCZOS)

    response = HttpResponse(content_type='image/png')
    img.save(response, 'PNG')
    return response


@require_GET
def device_qr_api(request, device_id):
    """
    API для получения QR-кода конкретного устройства.
    """
    device = get_object_or_404(Device, id=device_id)
    qr_code, created = QRCode.objects.get_or_create(device=device)
    filepath = os.path.join(settings.MEDIA_ROOT, 'qrcodes', f"{device.inventory_number}.png")
    if not os.path.exists(filepath):
        from ..qr_generator.utils import generate_qr_image_for_device
        generate_qr_image_for_device(device)
    return FileResponse(open(filepath, 'rb'), content_type='image/png')


@staff_member_required
def get_employees_by_department(request):
    """
    Возвращает JSON-список сотрудников, принадлежащих указанному отделу.
    Если department_id не передан, возвращает всех подтверждённых сотрудников.
    """
    department_id = request.GET.get('department_id')
    if department_id:
        employees = Employee.objects.filter(department_id=department_id, is_approved=True).values('id', 'full_name')
    else:
        employees = Employee.objects.filter(is_approved=True).values('id', 'full_name')
    return JsonResponse(list(employees), safe=False)


# ---------- Вспомогательные функции для импорта ----------
def get_or_create_department(name):
    if not name or str(name).strip() == '':
        return None
    name = str(name).strip()
    dept, _ = Department.objects.get_or_create(name=name)
    return dept


def get_or_create_employee(full_name, department=None):
    if not full_name or str(full_name).strip() == '':
        return None
    full_name = str(full_name).strip()
    try:
        emp = Employee.objects.get(full_name=full_name)
    except Employee.DoesNotExist:
        emp = Employee.objects.create(
            full_name=full_name,
            department=department,
            is_approved=False
        )
    return emp


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    elif isinstance(value, str):
        from django.utils.dateparse import parse_date
        d = parse_date(value)
        return d if d else None
    return None


def parse_status(status_str):
    """Преобразует строку статуса из Excel в значение для поля status модели."""
    if not status_str:
        return 'not_in_use'
    s = status_str.lower()
    if 'использование' in s:
        return 'in_use'
    elif 'резерв' in s:
        return 'reserve'
    else:
        return 'not_in_use'


def import_computers_from_sheet(sheet):
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    success = 0
    errors = []
    details = []
    for idx, row in enumerate(rows, start=2):
        try:
            inv = row[1]  # B
            if inv is None or str(inv).strip() == '':
                continue
            inv = str(int(inv)) if isinstance(inv, (int, float)) else str(inv).strip()

            name = str(row[2] or '').strip()  # C
            user_name = str(row[3] or '').strip()  # D
            dept_name = str(row[4] or '').strip()  # E

            department = get_or_create_department(dept_name)
            responsible = get_or_create_employee(user_name, department)

            os_info = str(row[6] or '').strip()  # G
            processor_brand = str(row[8] or '').strip()  # I
            processor_model = str(row[9] or '').strip()  # J
            ram_type = str(row[11] or '').strip()  # L
            ram_size = str(row[12] or '').strip()  # M
            disk_type = str(row[13] or '').strip()  # N
            ssd_size = str(row[14] or '').strip()  # O
            hdd_size = str(row[15] or '').strip()  # P
            purchase_date = parse_date(row[16])  # Q
            status_str = str(row[18] or '').lower()  # S
            status = parse_status(status_str)
            serial = str(row[21] or '').strip()  # V

            processor = f"{processor_brand} {processor_model}".strip()
            ram = ram_size
            disk_size = f"SSD:{ssd_size} HDD:{hdd_size}" if ssd_size or hdd_size else ''

            device_type, _ = DeviceType.objects.get_or_create(name='Компьютер', code='PC')

            computer, created = Computer.objects.update_or_create(
                inventory_number=inv,
                defaults={
                    'name': name,
                    'device_type': device_type,
                    'department': department,
                    'responsible': responsible,
                    'status': status,
                    'serial_number': serial,
                    'purchase_date': purchase_date,
                    'processor': processor,
                    'ram': ram,
                    'disk_size': disk_size,
                    'description': f"ОС: {os_info}\nТип RAM: {ram_type}\nТип накопителя: {disk_type}",
                }
            )
            action = 'создан' if created else 'обновлён'
            details.append(f"{action} компьютер {inv} – {computer.name}")
            success += 1
        except Exception as e:
            errors.append(f"Строка {idx}: {e}")
    return success, errors, details


def import_printers_from_sheet(sheet):
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    success = 0
    errors = []
    details = []
    for idx, row in enumerate(rows, start=2):
        try:
            inv = row[1]  # B
            if inv is None or str(inv).strip() == '':
                continue
            inv = str(int(inv)) if isinstance(inv, (int, float)) else str(inv).strip()

            name = str(row[2] or '').strip()  # C
            user_name = str(row[3] or '').strip()  # D
            dept_name = str(row[4] or '').strip()  # E
            paper_format = str(row[6] or '').strip()  # G
            purchase_date = parse_date(row[7])  # H
            status_str = str(row[10] or '').lower()  # K
            status = parse_status(status_str)
            serial = str(row[11] or '').strip()  # L

            department = get_or_create_department(dept_name)
            responsible = get_or_create_employee(user_name, department)

            device_type, _ = DeviceType.objects.get_or_create(name='Принтер', code='PRN')

            printer, created = Printer.objects.update_or_create(
                inventory_number=inv,
                defaults={
                    'name': name,
                    'device_type': device_type,
                    'department': department,
                    'responsible': responsible,
                    'status': status,
                    'serial_number': serial,
                    'purchase_date': purchase_date,
                    'color_type': 'bw',
                    'paper_format': paper_format,
                }
            )
            action = 'создан' if created else 'обновлён'
            details.append(f"{action} принтер {inv} – {printer.name}")
            success += 1
        except Exception as e:
            errors.append(f"Строка {idx}: {e}")
    return success, errors, details


# ---------- Представление для импорта (админка) ----------
@staff_member_required
@csrf_exempt
def import_devices_from_excel(request):
    if request.method == 'POST' and request.FILES.get('excel_file'):
        excel_file = request.FILES['excel_file']
        path = default_storage.save('tmp/' + excel_file.name, ContentFile(excel_file.read()))
        full_path = os.path.join(settings.MEDIA_ROOT, path)

        try:
            wb = openpyxl.load_workbook(full_path, data_only=True)
        except Exception as e:
            default_storage.delete(path)
            return JsonResponse({'status': 'error', 'message': f'Ошибка открытия файла: {e}'})

        results = {
            'computers': 0,
            'printers': 0,
            'errors': [],
            'details': []
        }

        if 'Гродно ПК' in wb.sheetnames:
            sheet = wb['Гродно ПК']
            comp_count, comp_errors, comp_details = import_computers_from_sheet(sheet)
            results['computers'] = comp_count
            results['errors'].extend(comp_errors)
            results['details'].extend(comp_details)

        if 'Гродно и бюро печатная техника' in wb.sheetnames:
            sheet = wb['Гродно и бюро печатная техника']
            prn_count, prn_errors, prn_details = import_printers_from_sheet(sheet)
            results['printers'] = prn_count
            results['errors'].extend(prn_errors)
            results['details'].extend(prn_details)

        default_storage.delete(path)
        return JsonResponse({'status': 'ok', 'results': results})

    return render(request, 'admin/import_excel.html')