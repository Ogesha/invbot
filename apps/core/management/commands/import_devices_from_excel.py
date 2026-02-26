import os
import openpyxl
from datetime import datetime
from django.core.management.base import BaseCommand, CommandError
from django.utils.dateparse import parse_date
from apps.core.models import Department, Employee, DeviceType, Computer, Printer

class Command(BaseCommand):
    help = 'Импорт устройств из Excel-файла (формат приложенной таблицы)'

    def add_arguments(self, parser):
        parser.add_argument('excel_file', type=str, help='Путь к Excel-файлу')

    def handle(self, *args, **options):
        file_path = options['excel_file']
        if not os.path.exists(file_path):
            raise CommandError(f'Файл {file_path} не найден')

        self.stdout.write(f'Загрузка файла {file_path}...')
        wb = openpyxl.load_workbook(file_path, data_only=True)  # data_only=True чтобы читать значения формул

        # Лист с компьютерами
        if 'Гродно ПК' not in wb.sheetnames:
            self.stdout.write(self.style.WARNING('Лист "Гродно ПК" не найден, пропускаем компьютеры'))
        else:
            sheet = wb['Гродно ПК']
            self.import_computers(sheet)

        # Лист с принтерами и МФУ
        if 'Гродно и бюро печатная техника' not in wb.sheetnames:
            self.stdout.write(self.style.WARNING('Лист "Гродно и бюро печатная техника" не найден, пропускаем принтеры'))
        else:
            sheet = wb['Гродно и бюро печатная техника']
            self.import_printers(sheet)

        self.stdout.write(self.style.SUCCESS('Импорт завершён'))

    def get_or_create_department(self, name):
        if not name or str(name).strip() == '':
            return None
        name = str(name).strip()
        dept, _ = Department.objects.get_or_create(name=name)
        return dept

    def get_or_create_employee(self, full_name, department=None):
        if not full_name or str(full_name).strip() == '':
            return None
        full_name = str(full_name).strip()
        # Пытаемся найти сотрудника с таким ФИО
        try:
            emp = Employee.objects.get(full_name=full_name)
        except Employee.DoesNotExist:
            emp = Employee.objects.create(
                full_name=full_name,
                department=department,
                is_approved=False  # по умолчанию неподтверждённый
            )
        return emp

    def parse_date(self, value):
        """Преобразует значение Excel в дату или None"""
        if isinstance(value, datetime):
            return value.date()
        elif isinstance(value, str):
            try:
                return parse_date(value).date()
            except:
                return None
        return None

    def import_computers(self, sheet):
        self.stdout.write('Импорт компьютеров...')
        # Предполагаем, что заголовки в первой строке, данные начинаются со второй
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        total = len(rows)
        success = 0
        errors = 0

        for idx, row in enumerate(rows, start=2):
            try:
                # Инвентарный номер (колонка B)
                inventory_number = row[1]
                if inventory_number is None or str(inventory_number).strip() == '':
                    self.stdout.write(self.style.WARNING(f'Строка {idx}: пропущена (нет инв. номера)'))
                    continue

                inventory_number = str(int(inventory_number)) if isinstance(inventory_number, (int, float)) else str(inventory_number).strip()

                # Наименование (колонка C)
                name = row[2] or ''
                name = str(name).strip()

                # Основной пользователь (колонка D)
                user_name = row[3] or ''
                # Структурное подразделение (колонка E)
                dept_name = row[4] or ''

                department = self.get_or_create_department(dept_name)
                responsible = self.get_or_create_employee(user_name, department)

                # Форм-фактор (F), ОС (G), Диагональ (H), Процессор (I), Модель процессора (J), Такт. частота (K)
                # Тип RAM (L), RAM (M), Тип накопителя (N), SSD (O), HDD (P)
                form_factor = str(row[5] or '').strip()
                os_info = str(row[6] or '').strip()
                diagonal = str(row[7] or '').strip()
                processor_brand = str(row[8] or '').strip()
                processor_model = str(row[9] or '').strip()
                processor_freq = str(row[10] or '').strip()
                ram_type = str(row[11] or '').strip()
                ram_size = str(row[12] or '').strip()  # CPU, ГБ – на самом деле RAM
                disk_type = str(row[13] or '').strip()
                ssd_size = str(row[14] or '').strip()
                hdd_size = str(row[15] or '').strip()

                # Дата ввода (Q)
                purchase_date = self.parse_date(row[16])

                # Статус (S) – столбец "Статус"
                status_str = str(row[18] or '').lower()
                status = 'использование' in status_str or 'резерв' in status_str

                # Серийный номер (V)
                serial_number = str(row[21] or '').strip()

                # Доп. оборудование (U), Антивирус (W), Новый домен (X)
                extra = str(row[20] or '').strip()
                antivirus = str(row[22] or '').strip()
                new_domain = str(row[23] or '').strip()

                # Собираем описание из всех оставшихся полей
                description_parts = []
                if form_factor:
                    description_parts.append(f"Форм-фактор: {form_factor}")
                if os_info:
                    description_parts.append(f"ОС: {os_info}")
                if diagonal:
                    description_parts.append(f"Диагональ: {diagonal}")
                if processor_brand and processor_model:
                    description_parts.append(f"Процессор: {processor_brand} {processor_model}")
                elif processor_brand:
                    description_parts.append(f"Процессор: {processor_brand}")
                if processor_freq:
                    description_parts.append(f"Тактовая частота: {processor_freq}")
                if ram_type:
                    description_parts.append(f"Тип RAM: {ram_type}")
                if disk_type:
                    description_parts.append(f"Тип накопителя: {disk_type}")
                if ssd_size and ssd_size != '0':
                    description_parts.append(f"SSD: {ssd_size} ГБ")
                if hdd_size and hdd_size != '0':
                    description_parts.append(f"HDD: {hdd_size} ГБ")
                if extra:
                    description_parts.append(f"Доп. оборудование: {extra}")
                if antivirus:
                    description_parts.append(f"Антивирус: {antivirus}")
                if new_domain:
                    description_parts.append(f"Новый домен: {new_domain}")

                description = '\n'.join(description_parts) if description_parts else ''

                # Поля для Computer
                processor = f"{processor_brand} {processor_model}".strip() if processor_brand or processor_model else ''
                ram = ram_size
                disk_size = f"SSD:{ssd_size} HDD:{hdd_size}" if ssd_size or hdd_size else ''

                # Получаем или создаём тип техники (должен быть создан ранее, но создадим при необходимости)
                device_type, _ = DeviceType.objects.get_or_create(name='Компьютер', code='PC')

                # Обновляем или создаём компьютер
                computer, created = Computer.objects.update_or_create(
                    inventory_number=inventory_number,
                    defaults={
                        'name': name,
                        'device_type': device_type,
                        'description': description,
                        'department': department,
                        'responsible': responsible,
                        'status': status,
                        'serial_number': serial_number,
                        'purchase_date': purchase_date,
                        'processor': processor,
                        'ram': ram,
                        'disk_size': disk_size,
                    }
                )
                if created:
                    self.stdout.write(f'  Создан компьютер {inventory_number}')
                else:
                    self.stdout.write(f'  Обновлён компьютер {inventory_number}')
                success += 1
            except Exception as e:
                self.stderr.write(self.style.ERROR(f'Строка {idx}: ошибка: {e}'))
                errors += 1

        self.stdout.write(f'Компьютеры: обработано {total}, успешно {success}, ошибок {errors}')

    def import_printers(self, sheet):
        self.stdout.write('Импорт принтеров и МФУ...')
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        total = len(rows)
        success = 0
        errors = 0

        for idx, row in enumerate(rows, start=2):
            try:
                # Инвентарный номер (колонка B)
                inventory_number = row[1]
                if inventory_number is None or str(inventory_number).strip() == '':
                    self.stdout.write(self.style.WARNING(f'Строка {idx}: пропущена (нет инв. номера)'))
                    continue
                inventory_number = str(int(inventory_number)) if isinstance(inventory_number, (int, float)) else str(inventory_number).strip()

                # Модель (C)
                name = row[2] or ''
                name = str(name).strip()

                # Основной пользователь (D)
                user_name = row[3] or ''
                # Структурное подразделение (E)
                dept_name = row[4] or ''

                department = self.get_or_create_department(dept_name)
                responsible = self.get_or_create_employee(user_name, department)

                # Форм-фактор (F)
                form_factor = str(row[5] or '').strip()
                # Формат печати (G)
                paper_format = str(row[6] or '').strip()
                # Дата ввода (H)
                purchase_date = self.parse_date(row[7])
                # Статус (K)
                status_str = str(row[10] or '').lower()
                status = 'использование' in status_str or 'резерв' in status_str
                # Серийный номер (L)
                serial_number = str(row[11] or '').strip()

                # Остальные поля (I, J, M) – формулы, игнорируем

                # Описание
                description_parts = []
                if form_factor:
                    description_parts.append(f"Форм-фактор: {form_factor}")
                description = '\n'.join(description_parts) if description_parts else ''

                # Тип техники
                device_type, _ = DeviceType.objects.get_or_create(name='Принтер', code='PRN')

                # Определяем цветность (из модели? в таблице нет, ставим 'bw')
                color_type = 'bw'

                printer, created = Printer.objects.update_or_create(
                    inventory_number=inventory_number,
                    defaults={
                        'name': name,
                        'device_type': device_type,
                        'description': description,
                        'department': department,
                        'responsible': responsible,
                        'status': status,
                        'serial_number': serial_number,
                        'purchase_date': purchase_date,
                        'color_type': color_type,
                        'paper_format': paper_format,
                    }
                )
                if created:
                    self.stdout.write(f'  Создан принтер {inventory_number}')
                else:
                    self.stdout.write(f'  Обновлён принтер {inventory_number}')
                success += 1
            except Exception as e:
                self.stderr.write(self.style.ERROR(f'Строка {idx}: ошибка: {e}'))
                errors += 1

        self.stdout.write(f'Принтеры: обработано {total}, успешно {success}, ошибок {errors}')