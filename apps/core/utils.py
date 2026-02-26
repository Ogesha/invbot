import openpyxl
from openpyxl.drawing.image import Image as XLImage
from django.http import HttpResponse
from datetime import datetime
import os
from PIL import Image
from io import BytesIO

# Импортируем синхронную функцию отправки логов из notifications
from ..bot.notifications import send_log_to_group_sync

def export_queryset_to_excel(queryset, filename_prefix, fields, headers):
    """
    Экспортирует queryset в Excel файл (только текст, без изображений).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Данные"

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = openpyxl.styles.Font(bold=True)

    for row_num, obj in enumerate(queryset, 2):
        for col_num, field in enumerate(fields, 1):
            value = obj
            parts = field.split('__')
            for part in parts:
                if value is None:
                    break
                value = getattr(value, part, None)
            if value is None:
                value = ''
            elif isinstance(value, datetime):
                value = value.strftime('%d.%m.%Y %H:%M')
            elif isinstance(value, bool):
                value = 'Да' if value else 'Нет'
            ws.cell(row=row_num, column=col_num).value = str(value)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f"{filename_prefix}_{timestamp}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def export_qrcodes_with_images_to_excel(queryset, filename_prefix='qrcodes'):
    """
    Экспортирует QR-коды в Excel с изображениями.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "QR-коды"

    headers = ['Изображение', 'Код QR', 'Устройство', 'Инвентарный номер']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = openpyxl.styles.Font(bold=True)

    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 30
    ws.column_dimensions['D'].width = 20

    row_num = 2
    for qr in queryset:
        code = qr.code
        device_name = qr.device.name if qr.device else '—'
        inventory_number = qr.device.inventory_number if qr.device else '—'

        ws.cell(row=row_num, column=2).value = code
        ws.cell(row=row_num, column=3).value = device_name
        ws.cell(row=row_num, column=4).value = inventory_number

        if qr.image and os.path.exists(qr.image.path):
            try:
                img = Image.open(qr.image.path)
                img.thumbnail((120, 120), Image.Resampling.LANCZOS)
                img_bytes = BytesIO()
                img.save(img_bytes, format='PNG')
                img_bytes.seek(0)
                xl_img = XLImage(img_bytes)
                xl_img.anchor = f'A{row_num}'
                ws.add_image(xl_img)
            except Exception as e:
                ws.cell(row=row_num, column=1).value = f"Ошибка: {str(e)[:30]}"
        else:
            ws.cell(row=row_num, column=1).value = "Нет изображения"

        ws.row_dimensions[row_num].height = 90
        row_num += 1

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f"{filename_prefix}_{timestamp}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response